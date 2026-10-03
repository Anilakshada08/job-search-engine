"""ATS resume rendering: .docx (python-docx) and .pdf (reportlab), plus a light attachment PDF.

ATS rules: single column, no tables/text boxes/icons, contact info in the body (not header/footer),
Calibri 10-11pt, US Letter, max 2 pages, headings Summary / Skills / Professional Experience / Education.
"""
from __future__ import annotations

import os
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from pypdf import PdfReader
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

from ..models import TailoredResume

FONT_DIRS = [Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts", Path("/usr/share/fonts/truetype/msttcorefonts"),
             Path("/Library/Fonts")]


def _register_calibri() -> tuple[str, str]:
    """Register Calibri for the ATS PDF; fall back to Helvetica if the font isn't installed."""
    for d in FONT_DIRS:
        reg, bold = d / "calibri.ttf", d / "calibrib.ttf"
        if reg.exists() and bold.exists():
            if "Calibri" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("Calibri", str(reg)))
                pdfmetrics.registerFont(TTFont("Calibri-Bold", str(bold)))
            return "Calibri", "Calibri-Bold"
    return "Helvetica", "Helvetica-Bold"


def _contact_line(profile: dict) -> str:
    return " | ".join(filter(None, [profile.get("home_location"), profile.get("phone"), profile.get("email"),
                                    profile.get("work_authorization")]))


def _role_line(e) -> str:
    who = e.company + (f" (Client: {e.client})" if e.client else "")
    return f"{e.title} | {who} | {e.dates}"


# -- DOCX ----------------------------------------------------------------------
def render_docx(r: TailoredResume, profile: dict, path: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(0.6))
    base = doc.styles["Normal"]
    base.font.name, base.font.size = "Calibri", Pt(10.5)
    base.paragraph_format.space_after = Pt(2)

    def para(text: str, size: float = 10.5, bold: bool = False, center: bool = False, style: str | None = None):
        p = doc.add_paragraph(style=style)
        run = p.add_run(text)
        run.font.name, run.font.size, run.bold = "Calibri", Pt(size), bold
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return p

    def heading(text: str):
        p = para(text.upper(), 11.5, bold=True)
        p.paragraph_format.space_before = Pt(8)

    para(profile["name"], 16, bold=True, center=True)
    para(r.headline, 11.5, bold=True, center=True)
    para(_contact_line(profile), 10, center=True)
    heading("Summary")
    para(r.summary)
    heading("Skills")
    para(", ".join(r.skills))
    heading("Professional Experience")
    for e in r.experience:
        para(_role_line(e), 10.5, bold=True).paragraph_format.space_before = Pt(5)
        if e.tech:
            para("Technologies: " + ", ".join(e.tech), 10)
        for b in e.bullets:
            para(b, style="List Bullet")
    heading("Education")
    for ed in r.education:
        para(ed)
    doc.save(str(path))


# -- PDF -----------------------------------------------------------------------
def _pdf(r: TailoredResume, profile: dict, path: Path, regular: str, bold: str) -> None:
    S = lambda name, **kw: ParagraphStyle(name, fontName=kw.pop("font", regular), fontSize=kw.pop("size", 10.5),
                                          leading=kw.pop("leading", 13), **kw)
    name_s = S("name", font=bold, size=16, leading=19, alignment=TA_CENTER)
    head_s = S("headline", font=bold, size=11.5, leading=14, alignment=TA_CENTER)
    contact_s = S("contact", size=10, alignment=TA_CENTER)
    sec_s = S("section", font=bold, size=11.5, leading=14, spaceBefore=8, spaceAfter=2)
    role_s = S("role", font=bold, spaceBefore=5)
    body_s = S("body")
    tech_s = S("tech", size=10)

    def esc(t: str) -> str:
        return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    story = [Paragraph(esc(profile["name"]), name_s), Paragraph(esc(r.headline), head_s),
             Paragraph(esc(_contact_line(profile)), contact_s),
             Paragraph("SUMMARY", sec_s), Paragraph(esc(r.summary), body_s),
             Paragraph("SKILLS", sec_s), Paragraph(esc(", ".join(r.skills)), body_s),
             Paragraph("PROFESSIONAL EXPERIENCE", sec_s)]
    for e in r.experience:
        story.append(Paragraph(esc(_role_line(e)), role_s))
        if e.tech:
            story.append(Paragraph(esc("Technologies: " + ", ".join(e.tech)), tech_s))
        story.append(ListFlowable([ListItem(Paragraph(esc(b), body_s), leftIndent=12) for b in e.bullets],
                                  bulletType="bullet", start="\u2022", leftIndent=12, bulletFontName=regular))
    story.append(Paragraph("EDUCATION", sec_s))
    story += [Paragraph(esc(ed), body_s) for ed in r.education]
    story.append(Spacer(1, 2))
    SimpleDocTemplate(str(path), pagesize=LETTER, leftMargin=0.6 * inch, rightMargin=0.6 * inch,
                      topMargin=0.6 * inch, bottomMargin=0.6 * inch,
                      title=f"{profile['name']} - {r.headline}", author=profile["name"]).build(story)


def render_pdf(r: TailoredResume, profile: dict, path: Path) -> None:
    _pdf(r, profile, path, *_register_calibri())


def render_attachment_pdf(r: TailoredResume, profile: dict, path: Path) -> None:
    """Same text, built-in Helvetica (no embedded fonts) so email attachments stay small."""
    _pdf(r, profile, path, "Helvetica", "Helvetica-Bold")


def pdf_pages(path: Path) -> int:
    return len(PdfReader(str(path)).pages)


def pdf_text(path: Path) -> str:
    text = " ".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    return re.sub(r"\s+", " ", text).lower()


def verify_keywords(path: Path, keywords: list[str]) -> list[str]:
    """Return the keywords that a text search of the rendered PDF cannot find."""
    text = pdf_text(path)
    squashed = text.replace(" ", "")
    missing = []
    for k in keywords:
        k_norm = re.sub(r"\s+", " ", k).lower().strip()
        if k_norm and k_norm not in text and k_norm.replace(" ", "") not in squashed:
            missing.append(k)
    return missing
